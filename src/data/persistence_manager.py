# -*- coding: utf-8 -*-
"""
HDF5波形数据持久化管理器
使用HDF5格式实现真正的数据追加，支持大文件高效处理
"""

import numpy as np
import h5py
import logging
import threading
import queue
import time
from pathlib import Path
from typing import Optional, Dict, List, Union
from datetime import datetime
import pandas as pd
import scipy.io
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)


class WaveformPersistence:
    """
    HDF5波形数据持久化管理器
    使用HDF5格式实现真正的数据追加，支持大文件高效处理
    """

    def __init__(self, base_dir: str = "data"):
        self.base_dir = Path(base_dir)
        self.base_dir.mkdir(parents=True, exist_ok=True)

        # 异步写入相关
        self._write_queue = queue.Queue(maxsize=5000)
        self._write_thread = None
        self._stop_event = threading.Event()
        self._current_file = None
        self._h5_file = None
        self._recording = False

        # 数据缓存 - 改为批量缓存所有通道
        self._data_cache = []  # list of channel_data_arrays
        self._cache_size = 0
        self._max_cache_size = 1000  # 缓存1000个样本后写入

    def start_recording(self, channel_config: List[Dict], sample_rate: int = 0, filename_prefix: Optional[str] = "", n_channels: int = 4) -> Path:
        """
        开始录制数据

        Args:
            filename_prefix: 自定义文件名前缀
            n_channels: 通道数量

        Returns:
            录制文件路径
        """
        if self._recording:
            logger.warning("已在录制中")
            return self._current_file

        # 生成文件名
        if filename_prefix == "":
            filename_prefix = "waveform"
        now = datetime.now()
        microseconds = now.strftime('%f')[:3]  # 取前3位微秒
        auto_name = f"{filename_prefix}_{now.strftime('%Y%m%d_%H%M%S')}{microseconds}.h5"
        self._current_file = self.base_dir / auto_name

        # 处理文件名冲突
        if self._current_file.exists():
            self._current_file = self._resolve_filename_conflict(self._current_file)

        # 创建HDF5文件并初始化数据集
        self._create_h5_file(n_channels, channel_config, sample_rate)

        # 启动写入线程
        self._start_write_thread()
        self._recording = True

        logger.info(f"开始录制波形数据到: {self._current_file}")
        return self._current_file

    def stop_recording(self):
        """停止录制数据"""
        if not self._recording:
            return

        self._recording = False

        # 刷新剩余缓存
        self._flush_cache()

        # 停止写入线程
        self._stop_write_thread()

        # 关闭HDF5文件
        if self._h5_file:
            self._h5_file.close()
            self._h5_file = None

        logger.info(f"停止录制，文件: {self._current_file}")
        self._current_file = None

    def _create_h5_file(self, n_channels: int, channel_config:List[Dict], sample_rate: int = 0):
        """创建HDF5文件并初始化数据集"""
        self._h5_file = h5py.File(self._current_file, 'w')

        # 创建元数据组
        metadata = self._h5_file.create_group('metadata')
        metadata.attrs['created_time'] = datetime.now().isoformat()
        metadata.attrs['version'] = '2.0.0'
        metadata.attrs['format'] = 'HDF5_Waveform'
        metadata.attrs['n_channels'] = n_channels
        metadata.attrs['sample_rate'] = sample_rate

        # 为每个通道创建可扩展数据集
        for ch in range(n_channels):
            # 获取通道配置
            ch_config = channel_config[ch]

            dataset = self._h5_file.create_dataset(
                f'channel_{ch:02d}',
                shape=(0,),  # 初始为空
                maxshape=(None,),  # 可无限扩展
                dtype=np.float32,
                chunks=True,  # 启用分块存储
                compression='gzip',  # 启用压缩
                compression_opts=1  # 压缩级别1（快速）
            )
            dataset.attrs['channel'] = ch
            dataset.attrs['unit'] = ch_config.get('unit', 'V')  # 从配置读取单位
            dataset.attrs['name'] = ch_config.get('name', f'CH{ch+1}')  # 通道名称
            # dataset.attrs['scale_factor'] = ch_config.get('scale_factor', 1.0)  # 缩放因子

        # 刷新到磁盘
        self._h5_file.flush()

    def push_save_data(self, channel_data: List[float]):
        """
        保存所有通道的数据到缓存

        Args:
            channel_data: 所有通道的数据列表，按通道顺序
        """
        if not self._recording:
            return

        # 添加到缓存
        self._data_cache.append(np.array(channel_data, dtype=np.float32))
        self._cache_size += 1

        # 检查是否需要刷新缓存
        if self._cache_size >= self._max_cache_size:
            self._flush_cache()

    def _flush_cache(self):
        """刷新缓存到写入队列"""
        if not self._data_cache:
            return

        try:
            # 将缓存数据转换为按通道分组的格式
            if self._data_cache:
                # 转换为numpy数组 (n_samples, n_channels)
                batch_data = np.array(self._data_cache)
                self._write_queue.put_nowait(batch_data)

            # 清空缓存
            self._data_cache.clear()
            self._cache_size = 0

        except queue.Full:
            logger.warning("写入队列已满，丢弃数据")

    def _start_write_thread(self):
        """启动写入线程"""
        if self._write_thread is not None:
            return

        self._stop_event.clear()
        self._write_thread = threading.Thread(
            target=self._write_worker,
            name="WaveformWriter",
            daemon=True
        )
        self._write_thread.start()

    def _stop_write_thread(self):
        """停止写入线程"""
        if self._write_thread is None:
            return

        try:
            self._stop_event.set()
            self._write_queue.put_nowait(None)  # 停止信号

            if self._write_thread.is_alive():
                self._write_thread.join(timeout=3.0)

            self._write_thread = None

        except Exception as e:
            logger.error(f"停止写入线程时出错: {e}")

    def _write_worker(self):
        """写入工作线程"""
        while not self._stop_event.is_set():
            try:
                # 获取数据
                try:
                    batch_data = self._write_queue.get(timeout=0.1)
                except queue.Empty:
                    continue

                if batch_data is None:  # 停止信号
                    break

                # 写入HDF5文件
                self._append_to_h5(batch_data)
                self._write_queue.task_done()

            except Exception as e:
                logger.error(f"写入线程错误: {e}")

        logger.info("波形写入线程已停止")

    def _append_to_h5(self, batch_data: np.ndarray):
        """追加数据到HDF5文件"""
        if not self._h5_file:
            return

        try:
            # batch_data shape: (n_samples, n_channels)
            n_samples, n_channels = batch_data.shape

            # 为每个通道追加数据
            for ch in range(n_channels):
                dataset_name = f'channel_{ch:02d}'
                if dataset_name in self._h5_file:
                    dataset = self._h5_file[dataset_name]

                    # 扩展数据集
                    old_size = dataset.shape[0]
                    new_size = old_size + n_samples
                    dataset.resize((new_size,))

                    # 追加新数据
                    dataset[old_size:new_size] = batch_data[:, ch]

            # 刷新到磁盘
            self._h5_file.flush()

        except Exception as e:
            logger.error(f"追加数据到HDF5文件失败: {e}")

    def _resolve_filename_conflict(self, filepath: Path) -> Path:
        """解决文件名冲突"""
        base_name = filepath.stem
        extension = filepath.suffix
        parent = filepath.parent

        counter = 1
        while True:
            new_name = f"{base_name}_{counter:03d}{extension}"
            new_path = parent / new_name
            if not new_path.exists():
                return new_path
            counter += 1

            if counter > 999:
                timestamp = int(time.time() * 1000) % 100000
                new_name = f"{base_name}_{timestamp}{extension}"
                return parent / new_name

    def load_data(self, filepath: Path) -> Dict[int, np.ndarray]:
        """
        从HDF5文件加载波形数据

        Args:
            filepath: 文件路径

        Returns:
            字典，键为通道索引，值为数据数组
        """
        try:
            data_dict = {}
            with h5py.File(filepath, 'r') as f:
                # 读取所有通道数据
                for key in f.keys():
                    if key.startswith('channel_'):
                        channel = int(key.replace('channel_', ''))
                        data_dict[channel] = f[key][:]  # 读取整个数据集

            logger.info(f"成功加载HDF5波形数据: {filepath}")
            return data_dict

        except Exception as e:
            logger.error(f"加载HDF5波形数据失败 {filepath}: {e}")
            return {}

    def is_recording(self) -> bool:
        """检查是否正在录制"""
        return self._recording

    def get_current_file(self) -> Optional[Path]:
        """获取当前录制文件路径"""
        return self._current_file

    def export_to_csv(self, hdf5_filepath: Path, output_filepath: Optional[Path] = None,
                     include_metadata: bool = True) -> Path:
        """
        将HDF5文件导出为CSV格式

        Args:
            hdf5_filepath: 源HDF5文件路径
            output_filepath: 输出CSV文件路径，如果为None则自动生成
            include_metadata: 是否包含元数据信息

        Returns:
            输出文件路径
        """
        try:
            # 生成输出文件路径
            if output_filepath is None:
                output_filepath = hdf5_filepath.with_suffix('.csv')

            # 加载数据
            data_dict = self.load_data(hdf5_filepath)
            if not data_dict:
                raise ValueError("无法加载HDF5数据")

            # 创建DataFrame
            df_data = {}
            max_length = 0

            # 获取所有通道数据
            for ch, data in data_dict.items():
                df_data[f'Channel_{ch:02d}'] = data
                max_length = max(max_length, len(data))

            # 创建时间轴
            with h5py.File(hdf5_filepath, 'r') as f:
                sample_rate = f['metadata'].attrs.get('sample_rate', 1000.0)

            time_axis = np.arange(max_length) / sample_rate
            df_data['Time_s'] = time_axis

            # 创建DataFrame
            df = pd.DataFrame(df_data)

            # 重新排列列顺序，时间轴在前
            cols = ['Time_s'] + [col for col in df.columns if col != 'Time_s']
            df = df[cols]

            # 导出CSV（暂时不包含元数据注释以避免编码问题）
            df.to_csv(output_filepath, index=False, encoding='utf-8-sig')

            # 如果需要包含元数据，创建单独的元数据文件
            if include_metadata:
                metadata_file = output_filepath.with_suffix('.metadata.txt')
                self._create_metadata_file(hdf5_filepath, metadata_file)

            logger.info(f"成功导出CSV文件: {output_filepath}")
            return output_filepath

        except Exception as e:
            logger.error(f"导出CSV失败: {e}")
            raise

    def export_to_excel(self, hdf5_filepath: Path, output_filepath: Optional[Path] = None,
                       include_metadata: bool = True) -> Path:
        """
        将HDF5文件导出为Excel格式

        Args:
            hdf5_filepath: 源HDF5文件路径
            output_filepath: 输出Excel文件路径，如果为None则自动生成
            include_metadata: 是否包含元数据工作表

        Returns:
            输出文件路径
        """
        try:
            # 生成输出文件路径
            if output_filepath is None:
                output_filepath = hdf5_filepath.with_suffix('.xlsx')

            # 加载数据
            data_dict = self.load_data(hdf5_filepath)
            if not data_dict:
                raise ValueError("无法加载HDF5数据")

            # 创建Excel工作簿
            wb = Workbook()

            # 删除默认工作表
            wb.remove(wb.active)

            # 创建数据工作表
            ws_data = wb.create_sheet("Data")

            # 准备数据
            df_data = {}
            max_length = 0

            for ch, data in data_dict.items():
                df_data[f'Channel_{ch:02d}'] = data
                max_length = max(max_length, len(data))

            # 创建时间轴
            with h5py.File(hdf5_filepath, 'r') as f:
                sample_rate = f['metadata'].attrs.get('sample_rate', 1000.0)

            time_axis = np.arange(max_length) / sample_rate
            df_data['Time_s'] = time_axis

            # 创建DataFrame
            df = pd.DataFrame(df_data)
            cols = ['Time_s'] + [col for col in df.columns if col != 'Time_s']
            df = df[cols]

            # 写入数据到Excel
            for r in dataframe_to_rows(df, index=False, header=True):
                ws_data.append(r)

            # 如果需要，添加元数据工作表
            if include_metadata:
                self._add_metadata_to_excel(hdf5_filepath, wb)

            # 保存Excel文件
            wb.save(output_filepath)

            logger.info(f"成功导出Excel文件: {output_filepath}")
            return output_filepath

        except Exception as e:
            logger.error(f"导出Excel失败: {e}")
            raise

    def export_to_matlab(self, hdf5_filepath: Path, output_filepath: Optional[Path] = None) -> Path:
        """
        将HDF5文件导出为MATLAB .mat格式

        Args:
            hdf5_filepath: 源HDF5文件路径
            output_filepath: 输出.mat文件路径，如果为None则自动生成

        Returns:
            输出文件路径
        """
        try:
            # 生成输出文件路径
            if output_filepath is None:
                output_filepath = hdf5_filepath.with_suffix('.mat')

            # 加载数据
            data_dict = self.load_data(hdf5_filepath)
            if not data_dict:
                raise ValueError("无法加载HDF5数据")

            # 准备MATLAB数据结构
            matlab_data = {}

            # 添加通道数据
            for ch, data in data_dict.items():
                matlab_data[f'channel_{ch:02d}'] = data

            # 添加元数据
            with h5py.File(hdf5_filepath, 'r') as f:
                metadata = dict(f['metadata'].attrs)
                matlab_data['metadata'] = metadata

                # 创建时间轴
                sample_rate = metadata.get('sample_rate', 1000.0)
                max_length = max(len(data) for data in data_dict.values())
                time_axis = np.arange(max_length) / sample_rate
                matlab_data['time_axis'] = time_axis

            # 保存为.mat文件
            scipy.io.savemat(output_filepath, matlab_data)

            logger.info(f"成功导出MATLAB文件: {output_filepath}")
            return output_filepath

        except Exception as e:
            logger.error(f"导出MATLAB失败: {e}")
            raise

    def _add_metadata_to_csv(self, hdf5_filepath: Path, csv_filepath: Path):
        """在CSV文件开头添加元数据注释"""
        try:
            # 读取现有CSV内容
            with open(csv_filepath, 'r', encoding='utf-8-sig') as f:
                csv_content = f.read()

            # 读取元数据
            with h5py.File(hdf5_filepath, 'r') as f:
                metadata = dict(f['metadata'].attrs)

            # 创建元数据注释
            metadata_lines = [
                "# HDF5 Waveform Data Export",
                f"# Source File: {hdf5_filepath.name}",
                f"# Export Time: {datetime.now().isoformat()}",
                "# Metadata:",
            ]

            for key, value in metadata.items():
                # 确保值是可序列化的字符串
                if isinstance(value, bytes):
                    value = value.decode('utf-8', errors='ignore')
                metadata_lines.append(f"# {key}: {value}")

            metadata_lines.append("# Data Format: Time_s, Channel_00, Channel_01, ...")
            metadata_lines.append("")  # 空行分隔

            # 重写文件
            with open(csv_filepath, 'w', encoding='utf-8-sig', newline='') as f:
                f.write('\n'.join(metadata_lines))
                f.write(csv_content)

        except Exception as e:
            logger.warning(f"添加CSV元数据失败: {e}")

    def _create_metadata_file(self, hdf5_filepath: Path, metadata_filepath: Path):
        """创建单独的元数据文件"""
        try:
            # 读取元数据
            with h5py.File(hdf5_filepath, 'r') as f:
                metadata = dict(f['metadata'].attrs)

            # 创建元数据文件
            with open(metadata_filepath, 'w', encoding='utf-8') as f:
                f.write("HDF5 Waveform Data Export Metadata\n")
                f.write("=" * 40 + "\n")
                f.write(f"Source File: {hdf5_filepath.name}\n")
                f.write(f"Export Time: {datetime.now().isoformat()}\n")
                f.write("\nMetadata:\n")
                f.write("-" * 20 + "\n")

                for key, value in metadata.items():
                    # 确保值是可序列化的字符串
                    if isinstance(value, bytes):
                        value = value.decode('utf-8', errors='ignore')
                    f.write(f"{key}: {value}\n")

                f.write("\nData Format:\n")
                f.write("-" * 20 + "\n")
                f.write("CSV columns: Time_s, Channel_00, Channel_01, Channel_02, ...\n")

        except Exception as e:
            logger.warning(f"创建元数据文件失败: {e}")

    def _add_metadata_to_excel(self, hdf5_filepath: Path, workbook: Workbook):
        """在Excel工作簿中添加元数据工作表"""
        try:
            # 创建元数据工作表
            ws_meta = workbook.create_sheet("Metadata", 0)  # 插入到第一个位置

            # 读取元数据
            with h5py.File(hdf5_filepath, 'r') as f:
                metadata = dict(f['metadata'].attrs)

            # 添加标题
            ws_meta.append(["HDF5 Waveform Data Export"])
            ws_meta.append([f"Source File: {hdf5_filepath.name}"])
            ws_meta.append([f"Export Time: {datetime.now().isoformat()}"])
            ws_meta.append([])  # 空行

            # 添加元数据
            ws_meta.append(["Metadata", "Value"])
            for key, value in metadata.items():
                ws_meta.append([key, str(value)])

        except Exception as e:
            logger.warning(f"添加Excel元数据失败: {e}")

    def export_batch(self, hdf5_filepath: Path, formats: List[str],
                    output_dir: Optional[Path] = None) -> Dict[str, Path]:
        """
        批量导出为多种格式

        Args:
            hdf5_filepath: 源HDF5文件路径
            formats: 要导出的格式列表 ['csv', 'excel', 'matlab']
            output_dir: 输出目录，如果为None则使用源文件目录

        Returns:
            格式到输出文件路径的映射
        """
        if output_dir is None:
            output_dir = hdf5_filepath.parent

        results = {}

        for fmt in formats:
            try:
                if fmt.lower() == 'csv':
                    output_path = output_dir / f"{hdf5_filepath.stem}.csv"
                    results['csv'] = self.export_to_csv(hdf5_filepath, output_path)

                elif fmt.lower() in ['excel', 'xlsx']:
                    output_path = output_dir / f"{hdf5_filepath.stem}.xlsx"
                    results['excel'] = self.export_to_excel(hdf5_filepath, output_path)

                elif fmt.lower() in ['matlab', 'mat']:
                    output_path = output_dir / f"{hdf5_filepath.stem}.mat"
                    results['matlab'] = self.export_to_matlab(hdf5_filepath, output_path)

                else:
                    logger.warning(f"不支持的导出格式: {fmt}")

            except Exception as e:
                logger.error(f"导出格式 {fmt} 失败: {e}")

        return results

    def __del__(self):
        """析构函数"""
        if self._recording:
            self.stop_recording()
